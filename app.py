import streamlit as st
from audio_recorder_streamlit import audio_recorder
from docx import Document
import io

st.set_page_config(page_title="Asistente de Contingencia", page_icon="🎙️", layout="centered")

st.title("🎙️ Asistente de Voz - Contingencia")
st.write("Graba tu audio, edita tu nota y genera el documento Word al tiro.")

# 1. Grabación
st.subheader("1. Grabación de Audio")
audio_bytes = audio_recorder(text="Haz clic para grabar", icon_size="2x")

if audio_bytes:
    st.audio(audio_bytes, format="audio/wav")
    st.success("¡Audio capturado con éxito!")

# 2. Texto
st.subheader("2. Registro de Notas")
texto_nota = st.text_area(
    "Escribe o edita el texto del reporte:", 
    placeholder="Escribe aquí los detalles...",
    height=150
)

# 3. Word
st.subheader("3. Descargar Word")

def generar_word(texto):
    doc = Document()
    doc.add_heading("Reporte de Asistente de Contingencia", 0)
    doc.add_paragraph(texto)
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

if st.button("Generar documento"):
    if texto_nota.strip() != "":
        archivo_docx = generar_word(texto_nota)
        st.success("¡Listo!")
        st.download_button(
            label="📥 Descargar archivo Word (.docx)",
            data=archivo_docx,
            file_name="registro_contingencia.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
    else:
        st.warning("Escribe algo antes de generar el documento.")
